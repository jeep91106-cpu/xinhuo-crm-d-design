# -*- coding: utf-8 -*-
import sys
from pathlib import Path

sys.path.insert(0, r"D:\CodexWorkspace\tools\python-docx")
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

src = Path(__file__).with_name("薪火CRM-重构PRD.md")
out = Path(__file__).with_name("薪火CRM-重构PRD.docx")
one = Path(__file__).parents[1] / "00-需求一页纸" / "需求一页纸.md"

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x1B, 0x16, 0x12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.25


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if size:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_cn(text, level):
    p = doc.add_paragraph()
    sizes = {0: 22, 1: 16, 2: 13, 3: 12}
    run = p.add_run(text)
    set_run_font(run, Pt(sizes.get(level, 12)), True, RGBColor(0x8F, 0x3D, 0x14) if level <= 1 else RGBColor(0x1B, 0x16, 0x12))
    p.paragraph_format.space_before = Pt(14 if level else 0)
    p.paragraph_format.space_after = Pt(8)
    return p


def flush_table(rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, Pt(9), bold=(i == 0))
    doc.add_paragraph()


def parse_md(text, title_override=None):
    rows = []
    first = True
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush_table(rows)
            rows = []
            continue
        if line.startswith("|") and "|" in line[1:]:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            rows.append(cells)
            continue
        flush_table(rows)
        rows = []
        if line.startswith("# "):
            add_heading_cn(title_override or line[2:].strip(), 0 if first else 1)
            first = False
            continue
        if line.startswith("## "):
            add_heading_cn(line[3:].strip(), 1)
            continue
        if line.startswith("### "):
            add_heading_cn(line[4:].strip(), 2)
            continue
        if line.startswith("```"):
            continue
        p = doc.add_paragraph()
        content = line
        if line.startswith("- "):
            content = "• " + line[2:]
        elif line[:2].isdigit() and line[1:3] == ". ":
            content = line
        run = p.add_run(content.replace("**", ""))
        set_run_font(run, Pt(11))
    flush_table(rows)


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("薪火 CRM 重构产品需求文档")
set_run_font(r, Pt(24), True, RGBColor(0x8F, 0x3D, 0x14))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("PRD v1.0  ·  2026-09-15  ·  配套需求一页纸与可编辑原型")
set_run_font(r, Pt(11), False, RGBColor(0x6E, 0x62, 0x56))

doc.add_paragraph()
add_heading_cn("0. 需求一页纸（已确认）", 1)
parse_md(one.read_text(encoding="utf-8"), title_override="（全文见仓库 00-需求一页纸）")

doc.add_page_break()
add_heading_cn("完整 PRD", 0)
# skip first H1 of PRD to avoid duplicate title
prd = src.read_text(encoding="utf-8")
prd_body = "\n".join(prd.splitlines()[1:])
parse_md(prd_body)

doc.save(out)
print(out)
