from pathlib import Path
import re,markdown
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
body=markdown.markdown(md,extensions=['tables','toc','fenced_code'])
from PIL import Image
def img_attrs(match):
    src=match.group(1); im=Image.open(ROOT/src); return '<img src="'+src+'" width="590" height="'+str(round(590*im.height/im.width))+'"'
body=re.sub(r'<img[^>]*src="([^"]+)"',img_attrs,body)
body=re.sub(r'(<tr>)(.*?</tr>)',lambda m:'<tr class="pending">'+m.group(2) if re.search(r'<td>(V0[1-8]|P0[1-4])</td>',m.group(2)) else m.group(0),body,flags=re.S)
css="""body{font-family:'Microsoft YaHei',Arial,sans-serif;color:#182b43;margin:0;background:#edf2f8;line-height:1.75}main{max-width:1120px;margin:24px auto;background:white;padding:52px 64px;box-shadow:0 4px 24px #17345818}h1{font-size:32px;color:#123964}h2{margin-top:60px;padding:12px 0;border-bottom:2px solid #2872b7;color:#164b7c}h3{margin-top:36px;color:#1b588b}p,li{font-size:15px}table{border-collapse:collapse;width:100%;font-size:13px;margin:22px 0;overflow-wrap:anywhere}th{background:#eaf2fc;text-align:left}th,td{border:1px solid #cdd9e6;padding:9px 11px;vertical-align:top}tr:nth-child(even){background:#f8fafc}a{color:#1c62a5}img{display:block;max-width:100%;width:900px;height:auto;margin:24px auto;border:1px solid #e4ebf3}tr.pending{color:#b5212f;background:#fff2f2}.toc{max-height:380px;overflow:auto;background:#f3f7fd;padding:18px;font-size:14px}.toc li{font-size:13px}strong{color:#123964}@media print{body{background:white}main{margin:0;padding:0;max-width:none;box-shadow:none}h2{break-before:page}h3,th{break-after:avoid}img{width:14cm;max-height:20cm;object-fit:contain;break-inside:avoid}tr{break-inside:avoid}thead{display:table-header-group}.toc{max-height:none;overflow:visible}table{font-size:9pt}p,li{font-size:10pt}@page{size:A4;margin:18mm 21mm}}"""
(ROOT/(STEM+'.html')).write_text('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+STEM+'</title><style>'+css+'</style></head><body><main>'+body+'</main></body></html>',encoding='utf-8')

doc=Document()
sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for sty in ['Normal','Title','Heading 1','Heading 2','Heading 3','List Bullet']:
    st=doc.styles[sty];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 2'].font.size=Pt(13)
head=sec.header.paragraphs[0];head.text='薪火 CRM · 完整产品需求文档 · 20260922修订V1.5 · 待评审';head.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('XH-CRM-PRD-001  ·  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def clean(s):
    s=re.sub(r'<[^>]+>','',s)
    s=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',s)
    return s.replace('**','').replace(chr(96),'')
def text_runs(p,s):
    # Minimal bold + actual hyperlink support; all content has one Markdown source.
    for token in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(<[^>]+>\)|\[[^\]]+\]\([^)]+\))',s):
        if token.startswith('**') and token.endswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('[') and '](' in token:
            label,url=token[1:].split('](',1);url=url[:-1].strip('<>')
            rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            hy=OxmlElement('w:hyperlink');hy.set(qn('r:id'),rel)
            run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1765A2');prop.append(color);run.append(prop)
            tx=OxmlElement('w:t');tx.text=label;run.append(tx);hy.append(run);p._p.append(hy)
        else:p.add_run(clean(token))
def make_table(rows):
    tb=doc.add_table(rows=1,cols=len(rows[0]));tb.style='Table Grid';tb.autofit=True
    for i,vals in enumerate(rows):
        cells=tb.rows[0].cells if i==0 else tb.add_row().cells
        for cell,v in zip(cells,vals):
            pa=cell.paragraphs[0];text_runs(pa,v);pa.paragraph_format.space_after=Pt(3)
            for r in pa.runs:
                r.font.size=Pt(8.5)
                if vals[0] in ['V01','V02','V03','V04','V05','V06','V07','V08','P01','P02','P03','P04']:r.font.color.rgb=RGBColor.from_string('B5212F')
                if i==0:r.bold=True
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF2FC');cell._tc.get_or_add_tcPr().append(shade)
        trp=tb.rows[i]._tr.get_or_add_trPr();nsp=OxmlElement('w:cantSplit');trp.append(nsp)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trp.append(repeat)
    doc.add_paragraph()
lines=md.splitlines();i=0;image_count=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='[TOC]':
        doc.add_heading('目录',1)
        p=doc.add_paragraph()
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='在 Word 中更新目录域可显示页码；网页版目录已可点击。';r.append(t);f.append(r);p._p.append(f)
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.match(r'^:?-+:?$',x) for x in cells):rows.append(cells)
            i+=1
        make_table(rows);continue
    im=re.match(r'^!\[(.*?)\]\((.*?)\)',line)
    if im:
        para=doc.add_paragraph();para.alignment=1;para.paragraph_format.keep_together=True
        para.add_run().add_picture(str(ROOT/im.group(2)),width=Cm(13.8));image_count+=1;i+=1;continue
    if line.startswith('# '):doc.add_heading(clean(line[2:]),0)
    elif line.startswith('## '):doc.add_heading(clean(line[3:]),1)
    elif line.startswith('### '):doc.add_heading(clean(line[4:]),2)
    elif line.startswith('- '):text_runs(doc.add_paragraph(style='List Bullet'),line[2:])
    else:
        para=doc.add_paragraph();text_runs(para,line)
        if line.startswith('**图 '):para.paragraph_format.keep_with_next=True
    i+=1
doc.core_properties.title=STEM;doc.core_properties.subject='内部广告业务与财务、人事行政完整重构 PRD';doc.core_properties.author='Codex';doc.core_properties.comments='待业务评审；文档检查不代表运行测试通过'
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(ROOT/(STEM+'.docx'))